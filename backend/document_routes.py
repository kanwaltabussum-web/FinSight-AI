"""FinSight AI supporting document inspection route.

This endpoint extracts readable text from PDF, DOCX, PPTX and TXT files. It
returns a transparent preview and a keyword-based type guess; it is not an AI
summarizer and does not perform transaction-level fraud/risk analysis.
"""
from io import BytesIO
from pathlib import Path
from datetime import datetime, timezone
import json
import re
from typing import Literal
from fastapi import APIRouter, File, UploadFile, HTTPException
from pydantic import BaseModel, Field

router = APIRouter()
MAX_DOCUMENT_BYTES = 20 * 1024 * 1024


def _classify(text: str, extension: str) -> tuple[str, str]:
    t = text.lower()
    if any(k in t for k in ("bank statement", "account statement", "opening balance", "closing balance", "iban", "transaction date")):
        return "Possible bank statement", "Text contains bank-statement terms. This is a keyword-based guess; check the original document."
    if any(k in t for k in ("invoice number", "invoice no", "bill to", "tax invoice", "amount due")):
        return "Possible invoice", "Text contains invoice-related terms. Confirm invoice details against the original file."
    if any(k in t for k in ("balance sheet", "profit and loss", "income statement", "cash flow statement", "total assets", "total liabilities")):
        return "Possible financial report", "Text contains financial-statement terms. This route does not calculate ratios or validate accounting balances."
    if any(k in t for k in ("audit report", "auditor's opinion", "internal control", "audit findings")):
        return "Possible audit document", "Text contains audit-related terms. This is not an audit opinion or compliance determination."
    if any(k in t for k in ("ifrs", "us gaap", "accounting standards", "regulatory update", "compliance policy")):
        return "Possible standards / compliance document", "Text contains standards or compliance terms. Verify the applicable jurisdiction and current requirements."
    if extension == ".pptx":
        return "Presentation", "Text was extracted from PowerPoint slides. Review the slide text preview for context."
    if extension == ".pdf":
        return "PDF document", "Readable text was extracted from the PDF. Scanned pages may require OCR."
    if extension == ".docx":
        return "Word document", "Text was extracted from the Word document, including table cell text where available."
    return "Text document", "Text was extracted from the uploaded file."


@router.post("/document/inspect")
async def inspect_document(file: UploadFile = File(...)):
    filename = file.filename or "uploaded_document"
    extension = Path(filename).suffix.lower()
    allowed = {".pdf", ".docx", ".pptx", ".txt"}
    if extension not in allowed:
        return {"success": False, "error": "Supported document formats are PDF, DOCX, PPTX and TXT. Legacy DOC/PPT files are not supported by this reader."}

    content = await file.read()
    if len(content) > MAX_DOCUMENT_BYTES:
        return {"success": False, "error": "This document is larger than the 20 MB limit."}
    if not content:
        return {"success": False, "error": "The uploaded document is empty."}

    text_parts = []
    page_count = slide_count = paragraph_count = None
    notes = []
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            page_count = len(reader.pages)
            for page in reader.pages:
                value = page.extract_text() or ""
                if value.strip(): text_parts.append(value.strip())
            if not text_parts:
                notes.append("No selectable text was found. The web client may attempt browser OCR on up to the first 10 pages if the OCR libraries can load.")
        elif extension == ".docx":
            from docx import Document
            doc = Document(BytesIO(content))
            paragraph_count = len(doc.paragraphs)
            text_parts.extend(p.text.strip() for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    line = " | ".join(c for c in cells if c)
                    if line: text_parts.append(line)
        elif extension == ".pptx":
            from pptx import Presentation
            presentation = Presentation(BytesIO(content))
            slide_count = len(presentation.slides)
            for number, slide in enumerate(presentation.slides, 1):
                slide_parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_parts.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            values = [cell.text.strip() for cell in row.cells]
                            line = " | ".join(v for v in values if v)
                            if line: slide_parts.append(line)
                if slide_parts:
                    text_parts.append(f"[Slide {number}]\n" + "\n".join(slide_parts))
        else:
            try:
                text_parts.append(content.decode("utf-8-sig"))
            except UnicodeDecodeError:
                text_parts.append(content.decode("latin-1", errors="replace"))
    except ImportError as exc:
        return {"success": False, "error": f"A required document-reading package is not installed: {exc.name}. Install backend/requirements.txt and redeploy the API."}
    except Exception as exc:
        return {"success": False, "error": f"Could not read this document: {str(exc)}"}

    extracted = "\n\n".join(text_parts).strip()
    words = re.findall(r"\b[\w'-]+\b", extracted)
    doc_type, explanation = _classify(extracted, extension)
    stats = {"file_size_kb": round(len(content) / 1024, 1), "word_count": len(words), "character_count": len(extracted)}
    if page_count is not None: stats["page_count"] = page_count
    if slide_count is not None: stats["slide_count"] = slide_count
    if paragraph_count is not None: stats["paragraph_count"] = paragraph_count

    return {
        "success": True,
        "filename": filename,
        "file_type": extension.lstrip("."),
        "detected_document_type": doc_type,
        "explanation": explanation,
        "word_count": len(words),
        "character_count": len(extracted),
        "page_count": page_count,
        "slide_count": slide_count,
        "paragraph_count": paragraph_count,
        "statistics": stats,
        "notes": notes,
        "text_preview": extracted[:8000],
        "text_truncated": len(extracted) > 8000,
        "processing_note": "Text extraction and keyword-based classification only; no transaction-level risk analysis or AI-generated summary was performed."
    }


class FeedbackPayload(BaseModel):
    name: str = Field(default="", max_length=80)
    contact_email: str = Field(default="", max_length=254)
    category: Literal["General feedback", "Bug report", "Feature request", "Financial analysis", "Usability", "Other"] = "General feedback"
    rating: int = Field(default=3, ge=1, le=5)
    message: str = Field(min_length=8, max_length=2500)
    website: str = Field(default="", max_length=250)  # honeypot: real visitors leave this blank


@router.post("/feedback")
async def submit_feedback(payload: FeedbackPayload):
    """Accept product feedback without exposing a public read endpoint."""
    if payload.website.strip():
        # Silently accept bot submissions, but do not store them.
        return {"success": True, "message": "Thank you for your feedback."}

    if payload.contact_email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", payload.contact_email):
        raise HTTPException(status_code=422, detail="Please provide a valid email address or leave the email field blank.")

    record = {
        "submitted_at": datetime.now(timezone.utc).isoformat(),
        "name": payload.name.strip(),
        "contact_email": payload.contact_email.strip(),
        "category": payload.category,
        "rating": payload.rating,
        "message": payload.message.strip(),
    }
    destination = Path(__file__).resolve().parent / "feedback_submissions.jsonl"
    try:
        with destination.open("a", encoding="utf-8") as feedback_file:
            feedback_file.write(json.dumps(record, ensure_ascii=False) + "\n")
    except OSError:
        raise HTTPException(status_code=503, detail="Feedback could not be saved by the server. Please try again later.")

    print(f"FinSight AI feedback received: category={record['category']!r}, rating={record['rating']}")
    return {"success": True, "message": "Thank you. Your feedback was received."}
